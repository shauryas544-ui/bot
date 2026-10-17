import os
import time
import unittest
from unittest.mock import patch, MagicMock
import docx
import ebooklib
from ebooklib import epub

import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from config import (
    TEMP_DIR,
    get_user_session,
    set_user_language,
    set_active_file,
    clear_active_file,
    get_valid_active_file,
    set_awaiting_split,
    sanitize_filename,
    cleanup_expired_temp_files,
)
from services.translator import (
    split_text_into_chunks,
    translate_text,
    translate_batch,
    translate_txt,
    translate_docx,
    translate_epub,
    translate_pdf,
    translate_file,
    _sanitize_epub_book,
    make_progress_bar,
    escape_markdown,
    format_progress_message,
    ProgressTracker,
)
from services.splitter import (
    split_txt,
    split_docx,
    split_epub,
    split_pdf,
    split_file,
)
from handlers.commands import build_language_keyboard


class TestConfigAndSession(unittest.TestCase):
    def test_session_management(self):
        user_id = 99999
        session = get_user_session(user_id)
        self.assertEqual(session.target_language, "hi")
        self.assertFalse(session.awaiting_split_count)
        self.assertIsNone(session.active_file)

        # Update language
        set_user_language(user_id, "fr")
        self.assertEqual(session.target_language, "fr")

        # Set active file
        dummy_file = os.path.join(TEMP_DIR, "dummy_test.txt")
        with open(dummy_file, "w") as f:
            f.write("hello")
        set_active_file(user_id, dummy_file, "dummy_test.txt", ".txt")
        self.assertIsNotNone(session.active_file)
        self.assertEqual(session.active_file.name, "dummy_test.txt")

        # Awaiting split
        set_awaiting_split(user_id, True)
        self.assertTrue(session.awaiting_split_count)

        # Clear active file and verify file deletion
        clear_active_file(user_id, remove_file=True)
        self.assertIsNone(session.active_file)
        self.assertFalse(session.awaiting_split_count)
        self.assertFalse(os.path.exists(dummy_file))

    def test_session_ttl_expiration(self):
        user_id = 88888
        dummy_file = os.path.join(TEMP_DIR, "dummy_expire.txt")
        with open(dummy_file, "w") as f:
            f.write("content")
        set_active_file(user_id, dummy_file, "dummy_expire.txt", ".txt")

        session = get_user_session(user_id)
        # Simulate elapsed time past TTL (e.g. 2 hours ago)
        session.active_file.created_at = time.time() - 7200
        self.assertTrue(session.active_file.is_expired(ttl_seconds=1800))

        # get_valid_active_file should detect expiration and prune
        valid_file = get_valid_active_file(user_id, ttl_seconds=1800)
        self.assertIsNone(valid_file)
        self.assertFalse(os.path.exists(dummy_file))

    def test_cleanup_expired_temp_files(self):
        # Create an old temp file
        old_file = os.path.join(TEMP_DIR, "temp_old_test.txt")
        with open(old_file, "w") as f:
            f.write("old data")
        # Backdate mtime by 2 hours
        two_hours_ago = time.time() - 7200
        os.utime(old_file, (two_hours_ago, two_hours_ago))

        # Create a fresh temp file
        fresh_file = os.path.join(TEMP_DIR, "temp_fresh_test.txt")
        with open(fresh_file, "w") as f:
            f.write("fresh data")

        try:
            cleaned = cleanup_expired_temp_files(max_age_seconds=1800)
            self.assertGreaterEqual(cleaned, 1)
            self.assertFalse(os.path.exists(old_file))
            self.assertTrue(os.path.exists(fresh_file))
        finally:
            if os.path.exists(fresh_file):
                os.remove(fresh_file)
            if os.path.exists(old_file):
                os.remove(old_file)

    def test_sanitize_filename(self):
        dirty = "../../etc/passwd:sample*file?.docx"
        clean = sanitize_filename(dirty)
        self.assertNotIn("/", clean)
        self.assertNotIn("\\", clean)
        self.assertNotIn(":", clean)
        self.assertTrue(clean.endswith(".docx"))

    def test_language_keyboard(self):
        markup = build_language_keyboard("hi")
        self.assertIsNotNone(markup.inline_keyboard)
        found_checked = False
        for row in markup.inline_keyboard:
            for btn in row:
                if "✅" in btn.text and btn.callback_data == "set_lang:hi":
                    found_checked = True
        self.assertTrue(found_checked)


def _mock_translate_chunk(text, target_lang="en", **kw):
    if "⟦§⟧" in text:
        parts = text.split("⟦§⟧")
        return " ⟦§⟧ ".join(f"[{target_lang}] {p.strip()}" for p in parts)
    return f"[{target_lang}] {text}"


@patch("services.translator._translate_single_chunk", side_effect=_mock_translate_chunk)
class TestTranslatorService(unittest.TestCase):
    def test_chunking_small_and_large(self, mock_translate):
        small = "Hello world! This is a simple sentence."
        chunks = split_text_into_chunks(small, max_chars=100)
        self.assertEqual(len(chunks), 1)

        # Large text with paragraphs
        paragraphs = ["Paragraph " + str(i) + " with some repeated text to fill buffer. " * 15 for i in range(20)]
        large = "\n\n".join(paragraphs)
        chunks = split_text_into_chunks(large, max_chars=500)
        self.assertGreater(len(chunks), 1)
        for c in chunks:
            self.assertLessEqual(len(c), 500)

    def test_translate_text(self, mock_translate):
        text = "Good morning, friend."
        res = translate_text(text, target_lang="hi")
        self.assertIn("[hi]", res)
        self.assertIn("Good morning", res)

    def test_translate_txt(self, mock_translate):
        input_file = os.path.join(TEMP_DIR, "sample_input.txt")
        output_file = os.path.join(TEMP_DIR, "sample_output.txt")
        with open(input_file, "w", encoding="utf-8") as f:
            f.write("Welcome to the automated translation system.\nEnjoy reading books.")

        try:
            res_path = translate_txt(input_file, output_file, target_lang="es")
            self.assertTrue(os.path.exists(res_path))
            with open(res_path, "r", encoding="utf-8") as f:
                content = f.read()
            self.assertIn("[es]", content)
        finally:
            if os.path.exists(input_file):
                os.remove(input_file)
            if os.path.exists(output_file):
                os.remove(output_file)

    def test_translate_docx(self, mock_translate):
        input_file = os.path.join(TEMP_DIR, "sample_input.docx")
        output_file = os.path.join(TEMP_DIR, "sample_output.docx")

        doc = docx.Document()
        doc.add_heading("Header One", 0)
        doc.add_paragraph("This is the first paragraph.")
        doc.save(input_file)

        try:
            res_path = translate_docx(input_file, output_file, target_lang="fr")
            self.assertTrue(os.path.exists(res_path))
            read_doc = docx.Document(res_path)
            self.assertTrue(len(read_doc.paragraphs) > 0)
            self.assertIn("[fr]", read_doc.paragraphs[1].text)
        finally:
            if os.path.exists(input_file):
                os.remove(input_file)
            if os.path.exists(output_file):
                os.remove(output_file)

    def test_translate_epub(self, mock_translate):
        input_file = os.path.join(TEMP_DIR, "sample_input.epub")
        output_file = os.path.join(TEMP_DIR, "sample_output.epub")

        book = epub.EpubBook()
        book.set_identifier("test-translate-epub")
        book.set_title("Translation Test")
        book.set_language("en")

        c1 = epub.EpubHtml(title="Chapter 1", file_name="c1.xhtml", lang="en")
        c1.content = "<html><body><h1>Title</h1><p>Hello world.</p></body></html>"
        book.add_item(c1)
        book.spine = ["nav", c1]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(input_file, book)

        try:
            res_path = translate_epub(input_file, output_file, target_lang="de")
            self.assertTrue(os.path.exists(res_path))
            read_b = epub.read_epub(res_path)
            self.assertEqual(read_b.get_metadata("DC", "language")[0][0], "de")
        finally:
            if os.path.exists(input_file):
                os.remove(input_file)
            if os.path.exists(output_file):
                os.remove(output_file)


    def test_translate_pdf(self, mock_translate):
        import pypdf
        input_file = os.path.join(TEMP_DIR, "sample_input.pdf")
        output_file = os.path.join(TEMP_DIR, "sample_output.docx")

        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=200, height=200)
        with open(input_file, "wb") as f:
            writer.write(f)

        try:
            res_path = translate_pdf(input_file, output_file, target_lang="es")
            self.assertTrue(os.path.exists(res_path))
            read_doc = docx.Document(res_path)
            self.assertTrue(len(read_doc.paragraphs) > 0)
        finally:
            if os.path.exists(input_file):
                os.remove(input_file)
            if os.path.exists(output_file):
                os.remove(output_file)


class TestSplitterService(unittest.TestCase):
    def test_split_txt(self):
        input_file = os.path.join(TEMP_DIR, "test_split.txt")
        with open(input_file, "w", encoding="utf-8") as f:
            for i in range(1, 11):
                f.write(f"Line {i}\n")

        parts = []
        try:
            parts = split_txt(input_file, lines_per_file=3, original_filename="test_split.txt")
            self.assertEqual(len(parts), 4)
            for p in parts:
                self.assertTrue(os.path.exists(p))
            with open(parts[0], "r", encoding="utf-8") as f:
                self.assertEqual(len(f.readlines()), 3)
            with open(parts[-1], "r", encoding="utf-8") as f:
                self.assertEqual(len(f.readlines()), 1)
        finally:
            if os.path.exists(input_file):
                os.remove(input_file)
            for p in parts:
                if os.path.exists(p):
                    os.remove(p)

    def test_split_docx(self):
        input_file = os.path.join(TEMP_DIR, "test_split.docx")
        doc = docx.Document()
        for i in range(1, 13):
            doc.add_paragraph(f"Paragraph {i}")
        doc.save(input_file)

        parts = []
        try:
            parts = split_docx(input_file, paragraphs_per_file=5, original_filename="test_split.docx")
            self.assertEqual(len(parts), 3)
            for p in parts:
                self.assertTrue(os.path.exists(p))
            d1 = docx.Document(parts[0])
            self.assertEqual(len(d1.paragraphs), 5)
            d3 = docx.Document(parts[2])
            self.assertEqual(len(d3.paragraphs), 2)
        finally:
            if os.path.exists(input_file):
                os.remove(input_file)
            for p in parts:
                if os.path.exists(p):
                    os.remove(p)

    def test_split_epub(self):
        input_file = os.path.join(TEMP_DIR, "test_split.epub")
        book = epub.EpubBook()
        book.set_identifier("test-split-epub")
        book.set_title("Splitting Test Book")
        book.set_language("en")

        chaps = []
        for i in range(1, 7):
            c = epub.EpubHtml(title=f"Chap {i}", file_name=f"ch_{i}.xhtml", lang="en")
            c.content = f"<html><body><p>Chap {i} body</p></body></html>"
            book.add_item(c)
            chaps.append(c)

        book.spine = ["nav"] + chaps
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(input_file, book)

        parts = []
        try:
            parts = split_epub(input_file, chapters_per_volume=2, original_filename="test_split.epub")
            self.assertEqual(len(parts), 3)
            for p in parts:
                self.assertTrue(os.path.exists(p))
                part_book = epub.read_epub(p)
                doc_items = [
                    item for item in part_book.get_items_of_type(ebooklib.ITEM_DOCUMENT)
                    if not item.file_name.endswith("nav.xhtml")
                ]
                self.assertEqual(len(doc_items), 2)
        finally:
            if os.path.exists(input_file):
                os.remove(input_file)
            for p in parts:
                if os.path.exists(p):
                    os.remove(p)

    def test_split_pdf(self):
        import pypdf
        input_file = os.path.join(TEMP_DIR, "test_split.pdf")
        writer = pypdf.PdfWriter()
        for _ in range(5):
            writer.add_blank_page(width=100, height=100)
        with open(input_file, "wb") as f:
            writer.write(f)

        parts = []
        try:
            parts = split_pdf(input_file, pages_per_file=2, original_filename="test_split.pdf")
            self.assertEqual(len(parts), 3)
            for p in parts:
                self.assertTrue(os.path.exists(p))
            r1 = pypdf.PdfReader(parts[0])
            self.assertEqual(len(r1.pages), 2)
            r3 = pypdf.PdfReader(parts[2])
            self.assertEqual(len(r3.pages), 1)
        finally:
            if os.path.exists(input_file):
                os.remove(input_file)
            for p in parts:
                if os.path.exists(p):
                    os.remove(p)

    def test_token_helpers(self):
        from config import mask_token, is_valid_token_format, verify_telegram_token, save_token_to_env
        from pathlib import Path
        from unittest.mock import MagicMock

        # Mask token
        self.assertEqual(mask_token(""), "<EMPTY>")
        self.assertTrue(mask_token("12345678:ABCDEFGHIJKLMN").startswith("12345678:********"))

        # Valid format
        self.assertTrue(is_valid_token_format("1234567890:ABCDEFGHIJKLMNOPQRSTUVWXYZ123456789"))
        self.assertFalse(is_valid_token_format("invalid_token"))
        self.assertFalse(is_valid_token_format(""))

        # Verify telegram token mock 200
        with patch("requests.get") as mock_get:
            mock_resp = MagicMock()
            mock_resp.status_code = 200
            mock_resp.json.return_value = {"ok": True, "result": {"username": "TestBot"}}
            mock_get.return_value = mock_resp
            valid, username, err = verify_telegram_token("1234567890:ABCDEFGHIJKLMNOPQRSTUVWXYZ123456789")
            self.assertTrue(valid)
            self.assertEqual(username, "TestBot")

        # Verify telegram token mock 401
        with patch("requests.get") as mock_get:
            mock_resp = MagicMock()
            mock_resp.status_code = 401
            mock_get.return_value = mock_resp
            valid, username, err = verify_telegram_token("1234567890:ABCDEFGHIJKLMNOPQRSTUVWXYZ123456789")
            self.assertFalse(valid)
            self.assertIn("401", err)

        # Save token to env
        test_env_path = Path(TEMP_DIR) / "test_env_output.env"
        try:
            save_token_to_env("99999999:TEST_SAVED_TOKEN_ABCD123456789012", test_env_path)
            content = test_env_path.read_text(encoding="utf-8")
            self.assertIn("TELEGRAM_BOT_TOKEN=99999999:TEST_SAVED_TOKEN_ABCD123456789012", content)
        finally:
            if test_env_path.exists():
                test_env_path.unlink()


class TestProgressTracking(unittest.TestCase):
    def test_make_progress_bar(self):
        bar_0 = make_progress_bar(0)
        self.assertEqual(bar_0, "░" * 10)

        bar_50 = make_progress_bar(50)
        self.assertEqual(bar_50, "▓" * 5 + "░" * 5)

        bar_100 = make_progress_bar(100)
        self.assertEqual(bar_100, "▓" * 10)

        # Clamping
        bar_neg = make_progress_bar(-10)
        self.assertEqual(bar_neg, "░" * 10)
        bar_over = make_progress_bar(150)
        self.assertEqual(bar_over, "▓" * 10)

    def test_escape_markdown(self):
        dirty = "Chapter_1 [Test]*bold* `code` (note)"
        clean = escape_markdown(dirty)
        for ch in ("*", "_", "`", "[", "]", "(", ")"):
            self.assertNotIn(ch, clean)

    def test_format_progress_message(self):
        msg = format_progress_message("Fortunately_I_Met_You.epub", "Hindi 🇮🇳", 50, "Chapter 6 of 12")
        self.assertIn("Fortunately_I_Met_You.epub", msg)
        self.assertIn("Hindi 🇮🇳", msg)
        self.assertIn("50%", msg)
        self.assertIn("Chapter 6 of 12", msg)
        self.assertIn("▓▓▓▓▓░░░░░", msg)

    def test_progress_tracker_throttling(self):
        bot_mock = MagicMock()
        loop_mock = MagicMock()
        loop_mock.is_closed.return_value = False

        tracker = ProgressTracker(
            bot=bot_mock,
            chat_id=12345,
            message_id=67890,
            filename="sample.epub",
            lang_name="Hindi 🇮🇳",
            loop=loop_mock,
        )

        with patch("asyncio.run_coroutine_threadsafe") as mock_dispatch:
            mock_dispatch.side_effect = lambda coro, loop: coro.close()
            # First update triggers
            tracker.update(5, "Discovered chapters")
            self.assertEqual(mock_dispatch.call_count, 1)

            # Rapid duplicate update within 2.0s without significant change is throttled
            tracker.update(6, "Discovered chapters")
            self.assertEqual(mock_dispatch.call_count, 1)

            # Completion (100%) always triggers immediately
            tracker.update(100, "Completed!")
            self.assertEqual(mock_dispatch.call_count, 2)

    @patch("services.translator._translate_single_chunk", side_effect=_mock_translate_chunk)
    def test_epub_with_progress_callback(self, mock_translate):
        input_file = os.path.join(TEMP_DIR, "sample_prog.epub")
        output_file = os.path.join(TEMP_DIR, "sample_prog_out.epub")

        book = epub.EpubBook()
        book.set_identifier("prog-test")
        book.set_title("Progress Book")
        book.set_language("en")

        c1 = epub.EpubHtml(title="Chapter 1", file_name="c1.xhtml", lang="en")
        c1.content = "<html><body><h1>Title 1</h1><p>Paragraph one.</p></body></html>"
        c2 = epub.EpubHtml(title="Chapter 2", file_name="c2.xhtml", lang="en")
        c2.content = "<html><body><h1>Title 2</h1><p>Paragraph two.</p></body></html>"
        book.add_item(c1)
        book.add_item(c2)
        book.spine = ["nav", c1, c2]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(input_file, book)

        progress_calls = []
        def callback(pct, detail):
            progress_calls.append((pct, detail))

        try:
            res = translate_epub(input_file, output_file, "hi", progress_callback=callback)
            self.assertTrue(os.path.exists(res))
            self.assertTrue(len(progress_calls) > 0)
            # Verify final call is 100%
            self.assertEqual(progress_calls[-1][0], 100)
            # Verify some chapter detail was recorded
            has_chapter_detail = any("Chapter" in call[1] or "section" in call[1].lower() for call in progress_calls)
            self.assertTrue(has_chapter_detail)
        finally:
            if os.path.exists(input_file):
                os.remove(input_file)
            if os.path.exists(output_file):
                os.remove(output_file)

    def test_sanitize_epub_book_none_attributes(self):
        """
        Tests that _sanitize_epub_book fixes NoneType attributes in metadata,
        items, links, and metas, preventing 'TypeError: Argument must be bytes or unicode, got NoneType'.
        """
        book = epub.EpubBook()
        book.IDENTIFIER_ID = None
        book.uid = None
        book.title = None

        # Add dirty metadata with None values and None attributes
        dc_ns = "http://purl.org/dc/elements/1.1/"
        book.metadata[dc_ns] = {
            "title": [(None, {"scheme": None, "id": None})],
            "creator": [(None, {"role": None})],
            "language": [(None, {})],
        }

        # Add item with None attributes
        item = epub.EpubHtml(title=None, file_name=None, lang=None)
        item.id = None
        item.media_type = None
        item.links = [{"rel": None, "href": None}]
        item.metas = [{"name": None, "content": None}]
        item.properties = [None]
        item.content = "<html><body><p>Cleaned</p></body></html>"
        book.add_item(item)
        book.spine = ["nav", item]

        # Add Link with uid=None into book.toc (the exact pattern produced by ebooklib's nav parser)
        book.toc = [epub.Link("item_1.xhtml", "Chapter 1", uid=None)]
        book.add_item(epub.EpubNcx())

        # Sanitize
        _sanitize_epub_book(book, target_lang="hi")

        self.assertIsNotNone(book.IDENTIFIER_ID)
        self.assertIsNotNone(book.uid)
        self.assertEqual(book.language, "hi")

        # Verify that write_epub succeeds without raising TypeError
        test_out = os.path.join(TEMP_DIR, "sanitized_test_out.epub")
        try:
            epub.write_epub(test_out, book)
            self.assertTrue(os.path.exists(test_out))
        finally:
            if os.path.exists(test_out):
                os.remove(test_out)

    @patch("services.translator._translate_single_chunk", side_effect=_mock_translate_chunk)
    def test_translate_batch(self, mock_translate):
        """
        Tests translate_batch with multiple inputs and ensures progress callback is called.
        """
        inputs = ["Hello world", "How are you doing?", "Telegram bot translation system"]
        batch_calls = []

        def on_batch(cur, tot):
            batch_calls.append((cur, tot))

        results = translate_batch(inputs, target_lang="es", progress_batch_callback=on_batch)
        self.assertEqual(len(results), 3)
        for r in results:
            self.assertIn("[es]", r)
        self.assertGreater(len(batch_calls), 0)


if __name__ == "__main__":
    unittest.main()


